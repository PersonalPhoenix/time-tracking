import json
import sqlite3
import os
from tkinter import messagebox
import time
import csv

from xlsxwriter import Workbook
from xlsxwriter.exceptions import FileCreateError
from docx import Document


none_export = 'Журнал пуст! Нечего экспортировать'

def get_dir_name() -> str:
    '''
    Получение дирекории для экспорта файла.
    '''

    try:
        with open('settings.json', 'r', encoding = 'utf-8') as file:
            dir_name: str = json.load(file)['dir_name']

            return dir_name
    except:
        messagebox.showerror('Внимание', 'Что-то случилось с settings.json\nФайл был сохранен в текущую директорию')
        dir_name: str = os.getcwd()

        return dir_name
    
def success() -> None:
    '''
    Уведомление об успехе.
    '''

    messagebox.showinfo('Успех', f'Файл был экспортирован в {get_dir_name()}')
    
def export_sql_to_excel() -> None:
    '''
    'Экспорт' -> 'Excel'.
    '''

    try:
        conn = sqlite3.connect('db.sqlite3')
        cursor = conn.cursor()

        cursor.execute('''select * from sessions;''')
        sql_query: list = cursor.fetchall()

        if sql_query == []:
            messagebox.showerror('Ошибка', f'{none_export}')
            return
        
        workbook = Workbook(f'{get_dir_name()}/экспорт от {time.strftime("%d-%m-%Y")} в {time.strftime("%H %M %S")}.xlsx')
        worksheet = workbook.add_worksheet()

        bold = workbook.add_format({'bold': True})
        worksheet.write('A1', '№', bold)
        worksheet.write('B1', 'Имя сессии', bold)
        worksheet.write('C1', 'Начал', bold)
        worksheet.write('D1', 'Закончил', bold)
        worksheet.write('E1', 'Рабочее время', bold)
        worksheet.write('F1', 'Пауза', bold)
        worksheet.write('G1', 'Начало паузы', bold)
        worksheet.write('H1', 'Конец паузы паузы', bold)
        worksheet.write('I1', 'Время паузы', bold)

        for i, row in enumerate(sql_query):
            worksheet.write(i+1, 0, row[0])
            worksheet.write(i+1, 1, row[1])
            worksheet.write(i+1, 2, row[2])
            worksheet.write(i+1, 3, row[3])
            worksheet.write(i+1, 4, row[4])
            worksheet.write(i+1, 5, row[5])
            worksheet.write(i+1, 6, row[6])
            worksheet.write(i+1, 7, row[7])
            worksheet.write(i+1, 8, row[8])
                        
        workbook.close()
        success()

    except FileCreateError:
        messagebox.showerror('ОШибка создания файла', 
                             f'Возника ошибка при создании файла.\nПроверьте корректность пути экспорта')

def export_sql_to_csv() -> None:
    '''
    'Экспорт' -> 'CSV'.
    '''

    try:
        conn = sqlite3.connect('db.sqlite3')
        cursor = conn.cursor()

        cursor.execute('''select * from sessions;''')
        data: list = cursor.fetchall()

        if data == []:
            messagebox.showerror('Ошибка', f'{none_export}')
            return

        with open(f'{get_dir_name()}/экспорт от {time.strftime("%d-%m-%Y")} в {time.strftime("%H %M %S")}.csv', 'w', newline = '') as file:
            writer = csv.writer(file)
            writer.writerow([i[0] for i in cursor.description])
            writer.writerows(data)

        cursor.close()
        conn.close()

        success()
        
    except:
        messagebox.showerror('ОШибка создания файла', 
                             f'Возника ошибка при создании файла.\nПроверьте корректность пути экспорта')

def export_sql_to_word() -> None:
    '''
    'Экспорт' -> 'Word'.
    '''
    
    try:
        conn = sqlite3.connect('db.sqlite3')
        cursor = conn.cursor()

        cursor.execute('''select * from sessions;''')
        data: list = cursor.fetchall()

        if data == []:
            messagebox.showerror('Ошибка', f'{none_export}')
            return
        
        doc = Document()

        headers = [description[0] for description in cursor.description]
        table = doc.add_table(rows=1, cols=len(headers))
        table.autofit = False
        table.style = 'Table Grid'
        header_cells = table.rows[0].cells
        for i, header in enumerate(headers):
            header_cells[i].text = header

        for row in data:
            cells = table.add_row().cells
            for i, column_value in enumerate(row):
                cells[i].text = str(column_value)

        doc.save(f'{get_dir_name()}/экспорт от {time.strftime("%d-%m-%Y")} в {time.strftime("%H %M %S")}.docx')

        cursor.close()
        conn.close()

        success()

    except:
        messagebox.showerror('ОШибка создания файла', 
                             f'Возника ошибка при создании файла\nПроверьте корректность пути экспорта')